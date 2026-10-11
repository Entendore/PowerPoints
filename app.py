import sys
import os
import tempfile
import subprocess
import shutil
from pathlib import Path
from datetime import datetime
import pymupdf  # replaces deprecated `import fitz`

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QFileDialog, QTextEdit, QFrame,
    QListWidget, QListWidgetItem, QProgressBar, QSplitter, QMessageBox,
    QButtonGroup, QToolButton, QSizePolicy
)
from PySide6.QtCore import (
    Qt, Signal, QThread, QTimer, QPointF, QPoint, QSize, QBuffer, QIODevice
)
from PySide6.QtGui import (
    QPixmap, QImage, QGuiApplication, QIcon,
    QDragEnterEvent, QDropEvent,
    QPainter, QPen, QColor, QPolygonF,
    QShortcut, QKeySequence
)
from pptx import Presentation


# ==========================================
# WORKER THREAD (Loads PPTX/HTML in Background)
# ==========================================
class LoaderThread(QThread):
    finished = Signal(list, list, list, str, str)  # images, notes, hidden, error, pdf_path

    def __init__(self, file_path):
        super().__init__()
        self.file_path = file_path

    @staticmethod
    def find_libreoffice():
        for cmd in ["libreoffice", "soffice", "soffice.exe", "libreoffice.exe"]:
            try:
                r = subprocess.run([cmd, "--version"], capture_output=True, timeout=10)
                if r.returncode == 0:
                    return cmd
            except (FileNotFoundError, subprocess.TimeoutExpired):
                continue
        return None

    def run(self):
        file_path = self.file_path
        images = []
        notes = []
        hidden_statuses = []

        # 1. Extract Notes & Hidden Status (Only works for PPTX)
        try:
            prs = Presentation(file_path)
            for slide in prs.slides:
                note = ""
                if slide.has_notes_slide:
                    note = slide.notes_slide.notes_text_frame.text
                notes.append(note)
                
                # Check if slide is hidden via XML
                is_hidden = 'show="0"' in slide._element.xml
                hidden_statuses.append(is_hidden)
        except Exception:
            # If it's an HTML file, this will fail, which is fine.
            pass

        # 2. Convert to PDF (Works for both PPTX and HTML via LibreOffice)
        lo_cmd = self.find_libreoffice()
        if not lo_cmd:
            self.finished.emit([], [], [], "LibreOffice not found. Please install it.", "")
            return

        temp_dir = tempfile.mkdtemp()
        try:
            subprocess.run(
                [lo_cmd, "--headless", "--convert-to", "pdf",
                 "--outdir", temp_dir, file_path],
                check=True, capture_output=True, timeout=180,
            )
        except Exception as e:
            self.finished.emit([], [], [], f"Conversion failed: {e}", "")
            return

        pdf_path = os.path.join(temp_dir, Path(file_path).stem + ".pdf")
        if not os.path.exists(pdf_path):
            self.finished.emit([], [], [], "PDF generation failed.", "")
            return

        # 3. Render to Images
        doc = pymupdf.open(pdf_path)
        zoom = 2.0
        mat = pymupdf.Matrix(zoom, zoom)
        for page in doc:
            pix = page.get_pixmap(matrix=mat)
            img_data = pix.tobytes("png")
            qimg = QImage()
            qimg.loadFromData(img_data, "PNG")
            images.append(qimg)
        doc.close()

        self.finished.emit(images, notes, hidden_statuses, "", pdf_path)


# ==========================================
# DRAWING CANVAS (Replaces QLabel)
# ==========================================
class DrawingCanvas(QWidget):
    def __init__(self):
        super().__init__()
        self.setMinimumSize(400, 300)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

        self.image = None
        self.scaled_pixmap = None
        self.scale_factor = 1.0
        self.offset = QPointF(0, 0)

        self.annotations = {}
        self.current_path = None
        self.current_index = -1

        self.tool = "none"
        self.color = Qt.red

    def set_image(self, img, index):
        self.image = img
        self.current_index = index
        self.update_scaling()
        self.update()

    def update_scaling(self):
        if not self.image:
            self.scaled_pixmap = None
            return
        pm = QPixmap.fromImage(self.image)
        sc = pm.scaled(self.size(), Qt.KeepAspectRatio, Qt.SmoothTransformation)
        self.scaled_pixmap = sc
        self.scale_factor = sc.width() / self.image.width()
        self.offset = QPointF((self.width() - sc.width()) / 2.0, (self.height() - sc.height()) / 2.0)

    def resizeEvent(self, event):
        self.update_scaling()
        super().resizeEvent(event)

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.fillRect(self.rect(), Qt.black)
        
        if not self.scaled_pixmap:
            return
            
        painter.drawPixmap(self.offset, self.scaled_pixmap)
        
        # Draw annotations using image coordinates mapped to screen
        painter.translate(self.offset)
        painter.scale(self.scale_factor, self.scale_factor)
        
        if self.current_index in self.annotations:
            for ann in self.annotations[self.current_index]:
                pen = QPen(ann["color"], 3.0)  # 3px thickness in image coords
                painter.setPen(pen)
                painter.drawPolyline(QPolygonF(ann["points"]))
                
        if self.current_path:
            pen = QPen(self.current_path["color"], 3.0)
            painter.setPen(pen)
            painter.drawPolyline(QPolygonF(self.current_path["points"]))

    def mousePressEvent(self, event):
        if self.tool == "pen" and event.button() == Qt.LeftButton:
            img_pos = (event.position() - self.offset) / self.scale_factor
            self.current_path = {"color": QColor(self.color), "points": [img_pos]}

    def mouseMoveEvent(self, event):
        if self.tool == "pen" and self.current_path:
            img_pos = (event.position() - self.offset) / self.scale_factor
            self.current_path["points"].append(img_pos)
            self.update()

    def mouseReleaseEvent(self, event):
        if self.tool == "pen" and self.current_path:
            if self.current_index not in self.annotations:
                self.annotations[self.current_index] = []
            self.annotations[self.current_index].append(self.current_path)
            self.current_path = None

    def set_tool(self, tool, color=Qt.red):
        self.tool = tool
        self.color = color
        if tool == "pen":
            self.setCursor(Qt.CrossCursor)
        else:
            self.unsetCursor()

    def clear_slide(self):
        if self.current_index in self.annotations:
            self.annotations[self.current_index] = []
            self.update()

    def erase_last(self):
        if self.current_index in self.annotations and self.annotations[self.current_index]:
            self.annotations[self.current_index].pop()
            self.update()


# ==========================================
# PRESENTER NOTES WINDOW
# ==========================================
class NotesWindow(QWidget):
    slide_changed = Signal(int)  # Emits absolute index

    def __init__(self):
        super().__init__()
        self.setWindowTitle("Presenter View")
        self.resize(1000, 700)
        self.images = []
        self.notes = []
        self.hidden = []
        self.current_index = 0
        self.start_time = None

        self.timer = QTimer(self)
        self.timer.timeout.connect(self.update_clocks)
        self.timer.start(1000)

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(8, 8, 8, 8)

        # Top Bar
        top_bar = QHBoxLayout()
        self.clock_label = QLabel("00:00:00")
        self.clock_label.setStyleSheet("font-size: 20px; font-weight: bold;")
        self.timer_label = QLabel("00:00:00")
        self.timer_label.setStyleSheet("font-size: 20px; font-weight: bold; color: #c0392b;")
        
        self.progress_bar = QProgressBar()
        self.progress_bar.setTextVisible(False)
        self.progress_bar.setFixedHeight(12)
        
        top_bar.addWidget(QLabel("Clock:"))
        top_bar.addWidget(self.clock_label)
        top_bar.addSpacing(20)
        top_bar.addWidget(QLabel("Elapsed:"))
        top_bar.addWidget(self.timer_label)
        top_bar.addStretch(1)
        top_bar.addWidget(self.progress_bar)
        main_layout.addLayout(top_bar)

        splitter = QSplitter(Qt.Horizontal)
        
        # Left: Thumbnails
        self.thumb_list = QListWidget()
        self.thumb_list.setIconSize(QSize(150, 110))
        self.thumb_list.setFlow(QListWidget.TopToBottom)
        self.thumb_list.setResizeMode(QListWidget.Adjust)
        self.thumb_list.setSpacing(5)
        self.thumb_list.currentRowChanged.connect(self.on_thumb_click)
        splitter.addWidget(self.thumb_list)

        # Right: Content
        right_panel = QWidget()
        right_layout = QVBoxLayout(right_panel)
        right_layout.setContentsMargins(0, 0, 0, 0)

        self.current_slide_label = QLabel("Current Slide")
        self.current_slide_label.setAlignment(Qt.AlignCenter)
        self.current_slide_label.setStyleSheet("background-color: black;")
        self.current_slide_label.setMinimumHeight(250)
        right_layout.addWidget(self.current_slide_label, 2)

        self.notes_edit = QTextEdit()
        self.notes_edit.setReadOnly(True)
        self.notes_edit.setStyleSheet("font-size: 16px; padding: 10px; background-color: #fffde7;")
        right_layout.addWidget(self.notes_edit, 3)

        next_layout = QHBoxLayout()
        self.next_slide_label = QLabel("Next Slide")
        self.next_slide_label.setAlignment(Qt.AlignCenter)
        self.next_slide_label.setStyleSheet("background-color: #2c3e50; color: white;")
        self.next_slide_label.setFixedHeight(120)
        self.next_slide_label.setFixedWidth(200)
        
        next_info_layout = QVBoxLayout()
        self.next_title = QLabel("Up Next:")
        self.next_title.setStyleSheet("font-weight: bold; font-size: 14px;")
        self.next_notes_label = QLabel("")
        self.next_notes_label.setWordWrap(True)
        self.next_notes_label.setStyleSheet("color: #555; font-size: 12px;")
        next_info_layout.addWidget(self.next_title)
        next_info_layout.addWidget(self.next_notes_label)
        next_info_layout.addStretch()
        
        next_layout.addWidget(self.next_slide_label)
        next_layout.addLayout(next_info_layout)
        right_layout.addLayout(next_layout, 1)

        splitter.addWidget(right_panel)
        splitter.setSizes([200, 800])
        main_layout.addWidget(splitter, 1)

        # Bottom Nav
        btn_layout = QHBoxLayout()
        self.prev_btn = QPushButton("◀ Previous")
        self.prev_btn.setStyleSheet("padding: 10px; font-size: 14px;")
        self.prev_btn.clicked.connect(lambda: self.slide_changed.emit(self.current_index - 1))
        
        self.next_btn = QPushButton("Next ▶")
        self.next_btn.setStyleSheet("padding: 10px; font-size: 14px;")
        self.next_btn.clicked.connect(lambda: self.slide_changed.emit(self.current_index + 1))
        
        btn_layout.addWidget(self.prev_btn)
        btn_layout.addStretch()
        self.slide_counter = QLabel("Slide 0 / 0")
        self.slide_counter.setStyleSheet("font-size: 16px; font-weight: bold;")
        btn_layout.addWidget(self.slide_counter)
        btn_layout.addStretch()
        btn_layout.addWidget(self.next_btn)
        main_layout.addLayout(btn_layout)

        # Setup keyboard shortcuts
        self.setup_shortcuts()

    def setup_shortcuts(self):
        """Register arrow-key shortcuts that work even when list/buttons have focus."""
        QShortcut(QKeySequence(Qt.Key_Right), self,
                  activated=lambda: self.slide_changed.emit(self.current_index + 1))
        QShortcut(QKeySequence(Qt.Key_Left), self,
                  activated=lambda: self.slide_changed.emit(self.current_index - 1))
        QShortcut(QKeySequence(Qt.Key_Up), self,
                  activated=lambda: self.slide_changed.emit(self.current_index - 1))
        QShortcut(QKeySequence(Qt.Key_Down), self,
                  activated=lambda: self.slide_changed.emit(self.current_index + 1))
        QShortcut(QKeySequence(Qt.Key_Space), self,
                  activated=lambda: self.slide_changed.emit(self.current_index + 1))
        QShortcut(QKeySequence(Qt.Key_PageDown), self,
                  activated=lambda: self.slide_changed.emit(self.current_index + 1))
        QShortcut(QKeySequence(Qt.Key_PageUp), self,
                  activated=lambda: self.slide_changed.emit(self.current_index - 1))

        # Prevent buttons/list from stealing focus
        for btn in [self.prev_btn, self.next_btn]:
            btn.setFocusPolicy(Qt.NoFocus)
        self.thumb_list.setFocusPolicy(Qt.NoFocus)

    def load_data(self, images, notes, hidden):
        self.images = images
        self.notes = notes
        self.hidden = hidden
        self.thumb_list.clear()
        
        for i, img in enumerate(images):
            pixmap = QPixmap.fromImage(img)
            scaled = pixmap.scaled(150, 110, Qt.KeepAspectRatio, Qt.SmoothTransformation)
            is_hidden = i < len(hidden) and hidden[i]
            item = QListWidgetItem(f"Slide {i+1}" + (" (Hidden)" if is_hidden else ""))
            item.setIcon(QIcon(scaled))
            if is_hidden:
                item.setForeground(Qt.red)
            self.thumb_list.addItem(item)
            
        self.start_time = datetime.now()
        self.update_slide(0)

    def on_thumb_click(self, row):
        if row >= 0:
            self.slide_changed.emit(row)

    def update_slide(self, index):
        if not self.images:
            return
        self.current_index = index
        self.thumb_list.blockSignals(True)
        self.thumb_list.setCurrentRow(index)
        self.thumb_list.blockSignals(False)
        
        pixmap = QPixmap.fromImage(self.images[index])
        self.current_slide_label.setPixmap(pixmap.scaled(
            self.current_slide_label.size(), Qt.KeepAspectRatio, Qt.SmoothTransformation
        ))
        
        note = self.notes[index].strip() if index < len(self.notes) else ""
        self.notes_edit.setPlainText(note if note else "(No speaker notes for this slide)")
        
        if index + 1 < len(self.images):
            next_pix = QPixmap.fromImage(self.images[index + 1])
            self.next_slide_label.setPixmap(next_pix.scaled(
                self.next_slide_label.size(), Qt.KeepAspectRatio, Qt.SmoothTransformation
            ))
            next_note = self.notes[index + 1].strip() if index + 1 < len(self.notes) else ""
            self.next_notes_label.setText(next_note[:100] + "..." if len(next_note) > 100 else next_note)
        else:
            self.next_slide_label.setText("End of Presentation")
            self.next_slide_label.setPixmap(QPixmap())
            self.next_notes_label.setText("")
            
        self.slide_counter.setText(f"Slide {index + 1} / {len(self.images)}")
        self.progress_bar.setMaximum(len(self.images))
        self.progress_bar.setValue(index + 1)

    def update_clocks(self):
        now = datetime.now()
        self.clock_label.setText(now.strftime("%H:%M:%S"))
        if self.start_time:
            elapsed = now - self.start_time
            secs = int(elapsed.total_seconds())
            hrs = secs // 3600
            mins = (secs % 3600) // 60
            ss = secs % 60
            self.timer_label.setText(f"{hrs:02d}:{mins:02d}:{ss:02d}")

    def keyPressEvent(self, event):
        # All navigation is handled by QShortcut now;
        super().keyPressEvent(event)


# ==========================================
# MAIN SLIDE VIEWER
# ==========================================
class SlideViewer(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("PPTX/HTML Viewer Ultimate")
        self.resize(1024, 640)
        self.setAcceptDrops(True)

        self.images = []
        self.notes = []
        self.hidden = []
        self.current_index = 0
        self.is_black = False
        self.loop_mode = False
        self.temp_pdf_path = ""
        
        self.loader = None

        central = QWidget()
        self.setCentralWidget(central)
        layout = QVBoxLayout(central)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        # Drawing Canvas
        self.canvas = DrawingCanvas()
        layout.addWidget(self.canvas, 1)

        # Control Bar
        self.control_bar = QFrame()
        self.control_bar.setStyleSheet("background-color: #2b2b2b;")
        cl = QHBoxLayout(self.control_bar)

        self.open_btn = QPushButton("📂 Open (Ctrl+O)")
        self.open_btn.setStyleSheet("color: white; padding: 8px;")
        self.open_btn.clicked.connect(self.open_file)

        self.pdf_btn = QPushButton("💾 Save PDF")
        self.pdf_btn.setStyleSheet("color: white; padding: 8px;")
        self.pdf_btn.clicked.connect(self.save_pdf)
        self.pdf_btn.setEnabled(False)

        self.html_btn = QPushButton("🌐 Export HTML")
        self.html_btn.setStyleSheet("color: white; padding: 8px;")
        self.html_btn.clicked.connect(self.export_html)
        self.html_btn.setEnabled(False)

        cl.addWidget(self.open_btn)
        cl.addWidget(self.pdf_btn)
        cl.addWidget(self.html_btn)
        
        # Drawing Tools
        cl.addWidget(QLabel("🔧 Tools:"))
        self.tool_group = QButtonGroup(self)
        self.tool_group.setExclusive(True)
        
        self.pen_none_btn = QToolButton()
        self.pen_none_btn.setText("Mouse")
        self.pen_none_btn.setCheckable(True)
        self.pen_none_btn.setChecked(True)
        self.pen_none_btn.clicked.connect(lambda: self.canvas.set_tool("none"))
        
        self.pen_red_btn = QToolButton()
        self.pen_red_btn.setText("🔴 Pen")
        self.pen_red_btn.setCheckable(True)
        self.pen_red_btn.clicked.connect(lambda: self.canvas.set_tool("pen", Qt.red))
        
        self.pen_blue_btn = QToolButton()
        self.pen_blue_btn.setText("🔵 Pen")
        self.pen_blue_btn.setCheckable(True)
        self.pen_blue_btn.clicked.connect(lambda: self.canvas.set_tool("pen", Qt.blue))
        
        self.erase_btn = QPushButton("↩ Erase")
        self.erase_btn.clicked.connect(self.canvas.erase_last)
        
        self.clear_btn = QPushButton("✖ Clear")
        self.clear_btn.clicked.connect(self.canvas.clear_slide)
        
        for btn in [self.pen_none_btn, self.pen_red_btn, self.pen_blue_btn, self.erase_btn, self.clear_btn]:
            btn.setStyleSheet("color: white; padding: 8px;")
            self.tool_group.addButton(btn)
            cl.addWidget(btn)

        cl.addStretch()
        
        self.loop_btn = QPushButton("🔁 Loop: OFF")
        self.loop_btn.setCheckable(True)
        self.loop_btn.setStyleSheet("color: white; padding: 8px;")
        self.loop_btn.clicked.connect(self.toggle_loop)
        
        self.prev_btn = QPushButton("◀ Prev")
        self.prev_btn.setStyleSheet("color: white; padding: 8px 15px;")
        self.prev_btn.clicked.connect(lambda: self.go_to_slide(self.current_index - 1))

        self.next_btn = QPushButton("Next ▶")
        self.next_btn.setStyleSheet("color: white; padding: 8px 15px;")
        self.next_btn.clicked.connect(lambda: self.go_to_slide(self.current_index + 1))

        self.black_btn = QPushButton("⬛ Black (B)")
        self.black_btn.setStyleSheet("color: white; padding: 8px;")
        self.black_btn.clicked.connect(self.toggle_black_screen)

        self.fullscreen_btn = QPushButton("⛶ Fullscreen (F)")
        self.fullscreen_btn.setStyleSheet("color: white; padding: 8px;")
        self.fullscreen_btn.clicked.connect(self.toggle_fullscreen)

        self.slide_info = QLabel("No file loaded")
        self.slide_info.setStyleSheet("color: white; padding: 8px;")

        cl.addWidget(self.loop_btn)
        cl.addWidget(self.prev_btn)
        cl.addWidget(self.slide_info)
        cl.addWidget(self.next_btn)
        cl.addWidget(self.black_btn)
        cl.addWidget(self.fullscreen_btn)
        layout.addWidget(self.control_bar)

        # Presenter Window
        self.notes_window = NotesWindow()
        self.notes_window.slide_changed.connect(self.go_to_slide)

        # Prevent buttons from stealing keyboard focus (so arrow keys always navigate)
        for btn in [
            self.open_btn, self.pdf_btn, self.html_btn,
            self.pen_none_btn, self.pen_red_btn, self.pen_blue_btn,
            self.erase_btn, self.clear_btn,
            self.loop_btn, self.prev_btn, self.next_btn,
            self.black_btn, self.fullscreen_btn,
        ]:
            btn.setFocusPolicy(Qt.NoFocus)

        # Setup keyboard shortcuts (works regardless of widget focus)
        self.setup_shortcuts()

    def setup_shortcuts(self):
        """Register global keyboard shortcuts that work even when buttons have focus."""

        # --- Arrow key navigation ---
        QShortcut(QKeySequence(Qt.Key_Right), self,
                  activated=lambda: self.go_to_slide(self.current_index + 1))
        QShortcut(QKeySequence(Qt.Key_Left), self,
                  activated=lambda: self.go_to_slide(self.current_index - 1))
        QShortcut(QKeySequence(Qt.Key_Up), self,
                  activated=lambda: self.go_to_slide(self.current_index - 1))
        QShortcut(QKeySequence(Qt.Key_Down), self,
                  activated=lambda: self.go_to_slide(self.current_index + 1))

        # --- Other navigation keys ---
        QShortcut(QKeySequence(Qt.Key_Space), self,
                  activated=lambda: self.go_to_slide(self.current_index + 1))
        QShortcut(QKeySequence(Qt.Key_PageDown), self,
                  activated=lambda: self.go_to_slide(self.current_index + 1))
        QShortcut(QKeySequence(Qt.Key_PageUp), self,
                  activated=lambda: self.go_to_slide(self.current_index - 1))
        QShortcut(QKeySequence(Qt.Key_Home), self,
                  activated=lambda: self.go_to_slide(0))
        QShortcut(QKeySequence(Qt.Key_End), self,
                  activated=lambda: self.go_to_slide(len(self.images) - 1))

        # --- View toggles ---
        QShortcut(QKeySequence(Qt.Key_F), self,
                  activated=self.toggle_fullscreen)
        QShortcut(QKeySequence(Qt.Key_B), self,
                  activated=self.toggle_black_screen)
        QShortcut(QKeySequence(Qt.Key_Escape), self,
                  activated=self._exit_fullscreen)

        # --- File / help ---
        QShortcut(QKeySequence(Qt.ControlModifier | Qt.Key_O), self,
                  activated=self.open_file)
        QShortcut(QKeySequence(Qt.Key_Question), self,
                  activated=self.show_shortcuts)
        QShortcut(QKeySequence(Qt.Key_F1), self,
                  activated=self.show_shortcuts)

    def _exit_fullscreen(self):
        if self.isFullScreen():
            self.showNormal()
            self.control_bar.show()

    # --- Navigation ---
    def go_to_slide(self, index):
        if not self.images:
            return
            
        # Handle Looping
        if self.loop_mode:
            if index < 0:
                index = len(self.images) - 1
            elif index >= len(self.images):
                index = 0
        else:
            if not (0 <= index < len(self.images)):
                return

        # Skip hidden slides during normal navigation
        while 0 <= index < len(self.images) and index < len(self.hidden) and self.hidden[index]:
            if index > self.current_index:
                index += 1
            else:
                index -= 1
            
            if not (0 <= index < len(self.images)):
                return

        self.current_index = index
        self.is_black = False
        self.show_slide()

    def show_slide(self):
        if not self.images:
            return
            
        if self.is_black:
            self.canvas.set_image(QImage(), -1) # Clear canvas to black
            self.canvas.setStyleSheet("background-color: black;")
        else:
            self.canvas.setStyleSheet("")
            self.canvas.set_image(self.images[self.current_index], self.current_index)
            
        self.slide_info.setText(f"Slide {self.current_index + 1} / {len(self.images)}")
        self.notes_window.update_slide(self.current_index)

    # --- File Handling ---
    def open_file(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Open Presentation", "", "Presentations (*.pptx *.ppt *.html *.htm);;All Files (*)"
        )
        if file_path:
            self.load_pptx(file_path)

    def dragEnterEvent(self, event: QDragEnterEvent):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def dropEvent(self, event: QDropEvent):
        for url in event.mimeData().urls():
            if url.toLocalFile().lower().endswith(('.pptx', '.ppt', '.html', '.htm')):
                self.load_pptx(url.toLocalFile())
                break

    def load_pptx(self, file_path):
        self.images = []
        self.canvas.set_image(QImage(), -1)
        self.canvas.update()
        self.slide_info.setText("Loading Presentation...")
        self.pdf_btn.setEnabled(False)
        self.html_btn.setEnabled(False)
        
        self.loader = LoaderThread(file_path)
        self.loader.finished.connect(self.on_loaded)
        self.loader.start()

    def on_loaded(self, images, notes, hidden, error, pdf_path):
        if error:
            QMessageBox.critical(self, "Error", error)
            return

        self.images = images
        self.notes = notes
        self.hidden = hidden
        self.current_index = 0
        self.temp_pdf_path = pdf_path
        self.pdf_btn.setEnabled(True)
        self.html_btn.setEnabled(True)
        
        self.show_slide()

        if self.images:
            self.notes_window.load_data(self.images, self.notes, self.hidden)
            self.position_notes_window()
            self.notes_window.show()
            self.notes_window.raise_()

    def save_pdf(self):
        if not self.temp_pdf_path or not os.path.exists(self.temp_pdf_path):
            return
        save_path, _ = QFileDialog.getSaveFileName(self, "Save PDF", "", "PDF Files (*.pdf)")
        if save_path:
            try:
                shutil.copy(self.temp_pdf_path, save_path)
                QMessageBox.information(self, "Success", "PDF saved successfully!")
            except Exception as e:
                QMessageBox.critical(self, "Error", f"Failed to save PDF:\n{e}")

    def export_html(self):
        if not self.images:
            QMessageBox.warning(self, "Warning", "No presentation loaded to export.")
            return
            
        save_path, _ = QFileDialog.getSaveFileName(self, "Export to HTML", "", "HTML Files (*.html)")
        if not save_path:
            return
            
        html_content = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>Presentation Export</title>
<style>
body { font-family: Arial, sans-serif; background: #f0f0f0; padding: 20px; margin: 0; }
h1 { color: #333; }
.slide-container { background: white; margin-bottom: 30px; border-radius: 8px; overflow: hidden; box-shadow: 0 4px 8px rgba(0,0,0,0.1); }
.slide-header { background: #2c3e50; color: white; padding: 10px; font-weight: bold; font-size: 18px; }
.slide-content { display: flex; flex-direction: row; }
.slide-img-container { width: 60%; background: #000; display: flex; justify-content: center; align-items: center; padding: 10px; }
.slide-img { max-width: 100%; max-height: 600px; object-fit: contain; }
.slide-notes { padding: 20px; width: 40%; white-space: pre-wrap; font-size: 16px; color: #444; overflow-y: auto; }
.hidden-tag { color: #e74c3c; font-size: 14px; margin-left: 10px; }
</style>
</head>
<body>
<h1>Presentation Export</h1>
"""
        for i, img in enumerate(self.images):
            buffer = QBuffer()
            buffer.open(QIODevice.WriteOnly)
            img.save(buffer, "PNG")
            img_base64 = buffer.data().toBase64().data().decode()
            
            note = self.notes[i].strip() if i < len(self.notes) else ""
            hidden_tag = '<span class="hidden-tag">(Hidden)</span>' if i < len(self.hidden) and self.hidden[i] else ""
            
            html_content += f"""
<div class="slide-container">
    <div class="slide-header">Slide {i+1} {hidden_tag}</div>
    <div class="slide-content">
        <div class="slide-img-container">
            <img class="slide-img" src="data:image/png;base64,{img_base64}" alt="Slide {i+1}">
        </div>
        <div class="slide-notes">{note if note else "(No speaker notes)"}</div>
    </div>
</div>
"""
        
        html_content += """
</body>
</html>
"""
        
        try:
            with open(save_path, "w", encoding="utf-8") as f:
                f.write(html_content)
            QMessageBox.information(self, "Success", "HTML exported successfully!")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to export HTML:\n{e}")

    def position_notes_window(self):
        screens = QGuiApplication.screens()
        if len(screens) > 1:
            main_screen = self.screen()
            for s in screens:
                if s is not main_screen:
                    geo = s.availableGeometry()
                    self.notes_window.move(geo.topLeft())
                    self.notes_window.resize(geo.width(), geo.height())
                    return
        self.notes_window.move(self.x() + self.width() + 20, self.y())

    # --- Toggles ---
    def toggle_loop(self):
        self.loop_mode = self.loop_btn.isChecked()
        self.loop_btn.setText("🔁 Loop: ON" if self.loop_mode else "🔁 Loop: OFF")

    def toggle_fullscreen(self):
        if self.isFullScreen():
            self.showNormal()
            self.control_bar.show()
        else:
            self.control_bar.hide()
            self.showFullScreen()

    def toggle_black_screen(self):
        if not self.images: return
        self.is_black = not self.is_black
        self.show_slide()

    def show_shortcuts(self):
        msg = QMessageBox(self)
        msg.setWindowTitle("Keyboard Shortcuts")
        msg.setText(
            "Navigation:\n"
            "  → / ↓ / Space / PgDn : Next slide\n"
            "  ← / ↑ / PgUp         : Previous slide\n"
            "  Home / End           : First / Last slide\n\n"
            "View:\n"
            "  F                    : Toggle Fullscreen\n"
            "  B                    : Toggle Black Screen\n"
            "  Esc                  : Exit Fullscreen\n\n"
            "Other:\n"
            "  Ctrl + O             : Open File\n"
            "  ? or F1              : Show this help"
        )
        msg.exec()

    # --- Keyboard ---
    def keyPressEvent(self, event):
        # All navigation is handled by QShortcut now;
        super().keyPressEvent(event)

    def closeEvent(self, event):
        # Cleanup temp directory on exit
        if self.temp_pdf_path and os.path.exists(os.path.dirname(self.temp_pdf_path)):
            try:
                shutil.rmtree(os.path.dirname(self.temp_pdf_path))
            except:
                pass
        event.accept()

def main():
    app = QApplication(sys.argv)
    app.setApplicationName("PPTX/HTML Viewer Ultimate")
    viewer = SlideViewer()
    viewer.show()

    if len(sys.argv) > 1 and os.path.exists(sys.argv[1]):
        QTimer.singleShot(100, lambda: viewer.load_pptx(sys.argv[1]))

    sys.exit(app.exec())

if __name__ == "__main__":
    main()